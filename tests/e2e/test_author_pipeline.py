"""One-book native service/artifact acceptance with explicit review boundaries.

Real DOCX parsing, canonical JS operations, agent evidence validation, image/
cover/QR rendering, EPUB/PDF preflight and exact-artifact packages. Provider
responses and human approvals are fixtures; no hosted Auth/DB/Storage/credit,
scanner, real provider quality/cost or retailer submission acceptance is implied.
"""
from __future__ import annotations

import base64
import copy
import hashlib
import json
import tempfile
import uuid
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from PIL import Image, ImageDraw
from pypdf import PdfReader

from test_critical_journey import _apply_operations, _client, WORKSPACE_ID, BOOK_MODEL


def _sha(data):
    return hashlib.sha256(data).hexdigest()


def _encoded(data):
    return base64.b64encode(data).decode()


def _job(monkeypatch, agent, model, chapter, tool, payload, approved_bible=()):
    import gateway
    provider = gateway.MockProvider([{"toolCalls": [{"name": tool, "input": payload}],
                                     "usage": {"inputTokens": 40, "outputTokens": 20}}])
    monkeypatch.setattr(gateway, "_REGISTRY", {"mock": lambda: provider})
    response = _client("ai").post("/v1/ai/jobs", json={
        "workspaceId": WORKSPACE_ID, "bookId": model["bookId"], "agentType": agent,
        "idempotencyKey": f"author-pipeline-{uuid.uuid4()}",
        "input": {"chapterIds": [chapter["id"]], "chapters": {chapter["id"]: chapter},
                  "book": model["metadata"], "bookBible": list(approved_bible)},
        "contextPolicy": {"includeBookBible": bool(approved_bible),
                          "includeStyleGuide": False, "includeRelatedContext": False,
                          "maxTokens": 4096},
    })
    assert response.status_code == 201, response.text
    assert len(provider.calls) == 1
    return response.json(), provider.calls[0]


def _operation(chapter, kind, payload, version):
    return {"operationId": str(uuid.uuid4()), "type": kind, "source": "human",
            "sourceRef": None, "target": {"chapterId": chapter["id"],
            "nodeId": payload.get("nodeId", payload.get("node", {}).get("id"))},
            "expectedVersion": version, "payload": payload}


def test_cross_language_operations_use_utf16_offsets_and_keep_original_text():
    model = copy.deepcopy(BOOK_MODEL)
    chapter = model["chapters"][0]
    chapter["nodes"][0]["text"] = "😀 The the lantern"
    original = copy.deepcopy(model)
    operation = _operation(chapter, "replace_text", {
        "nodeId": chapter["nodes"][0]["id"], "from": 3, "to": 10, "text": "The",
    }, 1)
    operation["source"] = "ai"
    result = _apply_operations(model, [operation], 1)
    assert result["bookModel"]["chapters"][0]["nodes"][0]["text"] == "😀 The lantern"
    assert result["version"] == 2 and model == original


@pytest.mark.parametrize("kind,bleed", [("ebook", None), ("print", "outer"), ("print", "all")])
def test_one_docx_book_reaches_reviewed_canon_artwork_and_native_packages(monkeypatch, kind, bleed):
    for name in ("DOCUMENT_SERVICE_TOKEN", "AI_SERVICE_TOKEN", "RENDERING_SERVICE_TOKEN", "PUBLISHING_SERVICE_TOKEN"):
        monkeypatch.setenv(name, "author-pipeline-private-fixture")
    monkeypatch.setenv("DEFAULT_AI_PROVIDER", "mock")
    monkeypatch.delenv("AI_RESULT_STORE", raising=False)
    headers = {"x-service-token": "author-pipeline-private-fixture"}

    original = Document()
    for number in range(1, 33):
        original.add_heading(f"Chapter {number}", level=1)
        paragraph = original.add_paragraph()
        if number == 1:
            paragraph.add_run("Mara reached the café. ").bold = True
            paragraph.add_run("The the lantern shone beside the harbor.").italic = True
        else:
            paragraph.add_run(f"Mara read harbor letter {number} and followed its clues.")
    source = BytesIO()
    original.save(source)
    source_bytes = source.getvalue()
    parsed = _client("document").post("/parse", headers=headers, json={
        "assetId": str(uuid.uuid4()), "format": "docx", "title": "Harbor Lantern",
        "contentBase64": _encoded(source_bytes),
    })
    assert parsed.status_code == 200, parsed.text
    model = parsed.json()["bookModel"]
    assert len(model["chapters"]) == 32
    model["metadata"].update(title="Harbor Lantern", author="Fixture Author", language="en")
    chapter = model["chapters"][0]
    node = next(node for node in chapter["nodes"] if node["type"] == "paragraph")
    assert node["attributes"]["richText"][0]["marks"] == [{"type": "bold"}]
    before_review = copy.deepcopy(model)
    start = node["text"].index("The the")
    # Operation offsets are UTF-16 units even though this client is Python.
    offset = len(node["text"][:start].encode("utf-16-le")) // 2
    edit = _operation(chapter, "replace_text", {"nodeId": node["id"], "from": offset,
                                              "to": offset + 7, "text": "The"}, 1)
    edit["source"] = "ai"
    reviewed, _ = _job(monkeypatch, "proofreader", model, {**chapter, "version": 1}, "propose_edit", {
        "chapterId": chapter["id"], "nodeId": node["id"], "operation": edit,
        "rationale": "Remove the doubled word without replacing the paragraph.", "confidence": 0.99,
    })
    assert reviewed["status"] == "succeeded", reviewed
    assert model == before_review, "generation must not modify the manuscript"
    suggestion_id = reviewed["suggestions"][0]["id"]
    accepted = _client("ai").post(f"/v1/ai/suggestions/{suggestion_id}/apply")
    assert accepted.status_code == 200, accepted.text
    applied = _apply_operations(model, [accepted.json()["operation"]], 1)
    assert applied["version"] == 2
    model = applied["bookModel"]
    assert _client("ai").post(f"/v1/ai/suggestions/{suggestion_id}/apply").status_code == 409
    chapter = model["chapters"][0]
    node = next(item for item in chapter["nodes"] if item["id"] == node["id"])
    assert "The the" not in node["text"] and "The lantern" in node["text"]
    assert node["attributes"]["richText"][0]["marks"] == [{"type": "bold"}]
    assert node["attributes"]["richText"][-1]["marks"] == [{"type": "italic"}]
    ref = {"chapterId": chapter["id"], "documentVersionId": str(uuid.uuid4()),
           "nodeId": node["id"], "textHash": _sha(node["text"].encode())}
    evidence = {**chapter, "version": 2, "documentVersionId": ref["documentVersionId"],
                "nodes": [{**item, "textHash": _sha(item.get("text", "").encode())} for item in chapter["nodes"]]}
    candidate = {"type": "character", "name": "Mara", "description": "Mara follows clues through harbor letters.",
                 "attributes": {"role": "protagonist"}, "sourceRefs": [ref], "confidence": 0.99}
    bible, _ = _job(monkeypatch, "bookbible", model, evidence, "propose_book_bible_candidates", {"candidates": [candidate]})
    assert bible["status"] == "succeeded", bible
    assert bible["suggestions"][0]["status"] == "pending" and model["bookBible"]["entities"] == []
    forged = {**candidate, "sourceRefs": [{**ref, "textHash": "0" * 64}]}
    rejected, _ = _job(monkeypatch, "bookbible", model, evidence, "propose_book_bible_candidates", {"candidates": [forged]})
    assert rejected["status"] == "failed" and rejected["suggestions"] == []
    approved = copy.deepcopy(bible["suggestions"][0])  # explicit fixture human approval, not a DB write
    model["bookBible"]["entities"] = [{"id": str(uuid.uuid4()), "type": approved["type"],
        "name": approved["name"], "description": approved["description"], "attributes": approved["attributes"],
        "sourceRefs": [ref["chapterId"]], "confidence": approved["confidence"]}]

    metadata = {"description": "Mara follows a lantern's clues through mysterious letters at a coastal harbor.",
                "keywords": ["harbor", "letters", "mystery"], "categories": ["FICTION / Mystery & Detective / General"],
                "audience": "Adult mystery readers", "rationale": "Based on the reviewed manuscript and approved Mara entry.",
                "confidence": 0.95, "sourceRefs": [ref]}
    identity = copy.deepcopy(model["metadata"])
    proposal, provider_call = _job(monkeypatch, "metadata", model, evidence, "propose_metadata", metadata, [approved])
    assert proposal["status"] == "succeeded", proposal
    assert model["metadata"] == identity and proposal["suggestions"][0]["status"] == "pending"
    prompt = json.dumps(provider_call["messages"])
    assert approved["description"] in prompt and ref["textHash"] in prompt
    foreign = {**metadata, "sourceRefs": [{**ref, "chapterId": str(uuid.uuid4())}]}
    rejected, _ = _job(monkeypatch, "metadata", model, evidence, "propose_metadata", foreign, [approved])
    assert rejected["status"] == "failed" and rejected["suggestions"] == []
    model["metadata"].update({key: proposal["suggestions"][0][key] for key in ("description", "keywords", "categories")})

    image_id, cover_id = str(uuid.uuid4()), str(uuid.uuid4())
    pixels = Image.new("RGB", (1600, 1000), "#123456")
    paint = ImageDraw.Draw(pixels)
    paint.rectangle((800, 0, 1599, 499), fill="#cc8844")
    paint.rectangle((0, 500, 799, 999), fill="#448866")
    image_data = BytesIO()
    pixels.save(image_data, "PNG")
    cover_data = BytesIO()
    Image.new("RGB", (1800, 2700), "#204050").save(cover_data, "PNG")
    model["assets"] = [{"id": image_id, "role": "illustration"}, {"id": cover_id, "role": "cover"}]
    placement = {"id": str(uuid.uuid4()), "type": "image", "assetId": image_id, "assetVersionNumber": 4,
                 "altText": "Mara's approved harbor lantern", "caption": "The harbor lantern"}
    placed = _apply_operations(model, [_operation(chapter, "insert_node", {
        "parentId": chapter["id"], "index": len(chapter["nodes"]), "node": placement,
    }, 2)], 2)
    model = placed["bookModel"]
    assert placed["version"] == 3 and model["chapters"][0]["nodes"][-1]["assetVersionNumber"] == 4
    destination = "https://author.example/harbor-lantern"
    edition = {"kind": kind, "cover": {"asset_id": cover_id, "qr_code": {
        "enabled": True, "url": destination, "size_px": 255, "position": "bottom-right"}},
        "front_matter": {"publisher": "Harbor Press", "copyright_notice": "Copyright Fixture Author. Permission required."}}
    if kind == "ebook":
        edition.update(navigation="toc+landmarks", include_title_page=True)
        channels = ("kdp", "apple", "barnesnoble", "googleplay")
    else:
        edition.update(trim_size="6x9", bleed_in=0.125, bleed_edges=bleed, include_table_of_contents=True,
            typography={"body_font": "BookwormVera", "heading_font": "BookwormVera-Bold"},
            wrap_cover={"enabled": True, "profile": "custom", "spine_width_in": 0.1, "expected_page_count": 36})
        channels = ("kdp", "barnesnoble") if bleed == "outer" else ("lulu",)
    request = {"bookModel": model, "editionConfig": edition, "coverBase64": _encoded(cover_data.getvalue()),
               "assetImagesBase64": {image_id: _encoded(image_data.getvalue())}}
    rendered = _client("rendering").post("/render", headers=headers, json=request)
    assert rendered.status_code == 200, rendered.text
    result = rendered.json()
    artifact = base64.b64decode(result["artifactBase64"], validate=True)
    assert _sha(artifact) == result["sha256"]
    artifact_directory = Path(tempfile.mkdtemp(prefix="bookworm-native-author-pipeline-"))
    extension = "epub" if kind == "ebook" else "pdf"
    (artifact_directory / f"book.{extension}").write_bytes(artifact)
    cover_artifact = base64.b64decode(result["coverArtifactBase64"], validate=True)
    assert _sha(cover_artifact) == result["coverSha256"]
    from cover_renderer import compose_front_cover, _qr_image
    from editions import parse_edition
    front_cover, _ = compose_front_cover(cover_data.getvalue(), model, parse_edition(edition))
    composed = Image.open(BytesIO(front_cover)).convert("RGB")
    pad, size = max(12, 255 // 12), 255
    right = composed.width - round(composed.width * 0.09) - pad
    bottom = round(composed.height * 0.9) - pad
    qr_pixels = composed.crop((right - size, bottom - size, right, bottom))
    assert qr_pixels.tobytes() == _qr_image(destination, size).tobytes()
    if kind == "ebook":
        assert cover_artifact == front_cover
        with zipfile.ZipFile(BytesIO(artifact)) as epub:
            opf = ET.fromstring(epub.read("OEBPS/content.opf"))
            assert opf.find(".//{http://purl.org/dc/elements/1.1/}description").text == metadata["description"]
            # Rich runs may wrap one sentence in several elements; compare semantic text.
            content = "\n".join("".join(ET.fromstring(epub.read(name)).itertext())
                                for name in epub.namelist() if name.endswith(".xhtml"))
            assert "The the lantern" not in content and "The lantern" in content and "Mara reached the café" in content
            chapter_dom = ET.fromstring(epub.read("OEBPS/ch0000.xhtml"))
            illustration = next(item for item in chapter_dom.iter() if item.tag.endswith("}img") and item.attrib.get("alt") == placement["altText"])
            embedded = Image.open(BytesIO(epub.read("OEBPS/" + illustration.attrib["src"]))).convert("RGB")
            assert embedded.tobytes() == pixels.tobytes()
        from epubcheck_runner import run_epubcheck
        checked = run_epubcheck(artifact)
        assert checked.status == "valid" and checked.errors == 0 and checked.warnings == 0, checked
    else:
        reader = PdfReader(BytesIO(artifact))
        assert len(reader.pages) == 36 and len(reader.outline) == 32
        text = "\n".join(page.extract_text() for page in reader.pages)
        assert "The lantern" in text and "The the lantern" not in text and "Mara reached the café" in text
        assert any(entry.image.convert("RGB").tobytes() == pixels.tobytes() for page in reader.pages for entry in page.images)
        assert any(entry.image.convert("RGB").tobytes() == composed.tobytes()
                   for page in PdfReader(BytesIO(cover_artifact)).pages for entry in page.images)
    for channel in channels:
        preflight = _client("rendering").post("/preflight", headers=headers, json={**request, "channel": channel})
        assert preflight.status_code == 200 and preflight.json()["errors"] == 0, preflight.text
        artifacts = {f"book.{extension}": result["artifactBase64"]}
        if kind == "print":
            artifacts["cover.pdf"] = result["coverArtifactBase64"]
        payload = {"bookModel": model, "editionConfig": edition, "channel": channel, "artifactsBase64": artifacts}
        exported = _client("publishing").post("/v1/publishing/package", headers=headers, json=payload)
        assert exported.status_code == 200, exported.text
        package = exported.json()["packages"][0]
        blob = base64.b64decode(package["dataBase64"], validate=True)
        assert _sha(blob) == package["sha256"]
        with zipfile.ZipFile(BytesIO(blob)) as archive:
            assert archive.read(f"book.{extension}") == artifact
            assert json.loads(archive.read("metadata.json"))["metadata"]["description"] == metadata["description"]
            assert json.loads(archive.read("manifest.json"))["channel"] == channel
            assert "NOT been submitted" in archive.read("README.txt").decode()
            if kind == "print":
                assert archive.read("cover.pdf") == cover_artifact
        replay = _client("publishing").post("/v1/publishing/package", headers=headers, json=payload)
        assert replay.json()["packages"][0]["sha256"] == package["sha256"]
        (artifact_directory / f"{channel}.zip").write_bytes(blob)
    (artifact_directory / "approved-book.json").write_text(json.dumps(model, ensure_ascii=False), encoding="utf-8")
    (artifact_directory / "front-cover.png").write_bytes(front_cover)
    assert _sha(source.getvalue()) == _sha(source_bytes), "original upload bytes changed"
    print(json.dumps({"stage": "native-author-pipeline", "kind": kind, "bleed": bleed,
                      "artifactSha256": result["sha256"], "channels": channels,
                      "path": str(artifact_directory), "provider": "mock", "submission": False}))
