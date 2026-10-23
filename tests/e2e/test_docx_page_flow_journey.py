"""DOCX -> authenticated parser -> real web serializer -> native EPUB/PDF bytes.

Synthetic manuscript, in-process service HTTP and native PDF inspection. This
does not assert hosted Storage, a browser session, or printer acceptance.
"""
import base64
import hashlib
import json
import os
import shutil
import subprocess
from io import BytesIO

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pypdf import PdfReader

from test_critical_journey import _client, REPO


@pytest.mark.parametrize("contents", [False, True])
def test_imported_page_flow_survives_editor_save_epub_reimport_and_pdf(monkeypatch, tmp_path, contents):
    monkeypatch.setenv("DOCUMENT_SERVICE_TOKEN", "page-flow-fixture-only")
    monkeypatch.setenv("RENDERING_SERVICE_TOKEN", "page-flow-fixture-only")
    headers = {"x-service-token": "page-flow-fixture-only"}
    source = Document()
    style = source.styles.add_style("Author chapter", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = source.styles["Heading 1"]
    style.paragraph_format.page_break_before = True
    source.add_paragraph("Opening chapter", style=style)
    run = source.add_paragraph().add_run("BEFORE MANUAL")
    run.bold = True
    run.add_break(WD_BREAK.PAGE)
    run.add_text("AFTER MANUAL")
    source.add_page_break()
    source.add_page_break()  # The author's deliberate blank page must remain.
    source.add_paragraph("BEYOND BLANK")
    source.add_page_break()  # A chapter boundary must not add a second break.
    source.add_paragraph("Closing chapter", style=style)
    source.add_paragraph("FINAL BODY")
    buffer = BytesIO()
    source.save(buffer)
    raw = buffer.getvalue()
    source_hash = hashlib.sha256(raw).hexdigest()
    parsed = _client("document").post("/parse", headers=headers, json={
        "assetId": "page-flow-docx", "format": "docx", "title": "Page Flow",
        "contentBase64": base64.b64encode(raw).decode(),
    })
    assert parsed.status_code == 200, parsed.text
    model = parsed.json()["bookModel"]
    assert len(model["chapters"]) == 2
    assert sum(n["type"] == "pageBreak" for ch in model["chapters"] for n in ch["nodes"]) == 4
    executable = shutil.which("node")
    assert executable, "Node is required for the real web editor serialization proof"
    environment = {k: os.environ[k] for k in ("PATH", "SystemRoot", "WINDIR", "TEMP", "TMP") if k in os.environ}
    environment["NODE_ENV"] = "test"
    edited = subprocess.run([executable, "--import", "tsx", "tests/e2e/docx-page-flow-editor.mjs"],
        cwd=REPO, input=json.dumps(model), text=True, encoding="utf-8", capture_output=True,
        timeout=30, env=environment)
    assert edited.returncode == 0, edited.stderr
    model = json.loads(edited.stdout)

    artifacts = {}
    for kind in ("ebook", "print"):
        config = {"kind": kind, **({"include_table_of_contents": contents} if kind == "print" else {})}
        request = {"bookModel": model, "editionConfig": config}
        response = _client("rendering").post("/render", headers=headers, json=request)
        assert response.status_code == 200, response.text
        data = response.json()
        blob = base64.b64decode(data["artifactBase64"], validate=True)
        assert hashlib.sha256(blob).hexdigest() == data["sha256"]
        repeated = _client("rendering").post("/render", headers=headers, json=request)
        assert repeated.status_code == 200 and repeated.json()["sha256"] == data["sha256"]
        artifacts[kind] = blob

    pages = [page.extract_text() for page in PdfReader(BytesIO(artifacts["print"])).pages]
    def page_of(marker):
        indices = [i for i, text in enumerate(pages) if marker in text]
        assert len(indices) == 1, f"Missing or duplicated fixture marker: {marker}"
        return indices[0]
    assert page_of("AFTER MANUAL Edited.") == page_of("BEFORE MANUAL") + 1
    assert page_of("BEYOND BLANK") == page_of("AFTER MANUAL Edited.") + 2
    assert page_of("FINAL BODY") == page_of("BEYOND BLANK") + 1
    assert "Opening chapter" in pages[page_of("BEFORE MANUAL")]
    assert "Closing chapter" in pages[page_of("FINAL BODY")]
    for title in ("Opening chapter", "Closing chapter"):
        assert sum(text.count(title) for text in pages) == 1 + int(contents)

    reimport = _client("document").post("/parse", headers=headers, json={
        "assetId": "page-flow-epub", "format": "epub", "title": "Page Flow",
        "contentBase64": base64.b64encode(artifacts["ebook"]).decode(),
    })
    assert reimport.status_code == 200, reimport.text
    def semantic(book):
        return [(ch["title"], [(n["type"], n.get("text"), n.get("level")) for n in ch["nodes"]]) for ch in book["chapters"]]
    assert semantic(reimport.json()["bookModel"]) == semantic(model)
    assert hashlib.sha256(raw).hexdigest() == source_hash
    from epubcheck_runner import run_epubcheck
    checked = run_epubcheck(artifacts["ebook"])
    assert checked.status == "valid" and checked.errors == 0 and checked.warnings == 0, checked
    (tmp_path / "source.docx").write_bytes(raw)
    (tmp_path / "book.pdf").write_bytes(artifacts["print"])
    (tmp_path / "book.epub").write_bytes(artifacts["ebook"])
    receipt = {"status": "passed", "synthetic": True, "contents": contents,
        "pdfRenderer": "pdf-1.18.0", "pages": len(pages), "pageBreaks": 4,
        "pdfSha256": hashlib.sha256(artifacts["print"]).hexdigest(),
        "epubSha256": hashlib.sha256(artifacts["ebook"]).hexdigest(),
        "sourceSha256": source_hash, "epubcheck": checked.status,
        "proofDirectory": str(tmp_path)}
    (tmp_path / "receipt.json").write_text(json.dumps(receipt, indent=2), encoding="utf-8")
    print(json.dumps(receipt))
