"""Real Word archives preserve deliberate page flow and outline semantics."""
import io
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from parsers import ParseError
from parsers.docx_parser import parse_docx


def archive(document):
    result = io.BytesIO()
    document.save(result)
    return result.getvalue()


def all_nodes(book):
    return [n for chapter in book["chapters"] for n in chapter["nodes"]]


def outline(properties, value):
    element = OxmlElement("w:outlineLvl")
    element.set(qn("w:val"), str(value))
    properties.append(element)


def test_intra_run_page_break_keeps_text_marks_tabs_and_soft_breaks_once():
    doc = Document()
    run = doc.add_paragraph().add_run("Before\tpage\nline")
    run.bold = True
    run.add_break(WD_BREAK.PAGE)
    # Word's cached pagination marker must not become a second authored break.
    run._r.append(OxmlElement("w:lastRenderedPageBreak"))
    run.add_text("After page")
    book, _ = parse_docx(archive(doc))
    nodes = all_nodes(book)
    assert [n["type"] for n in nodes] == ["paragraph", "pageBreak", "paragraph"]
    assert [n["text"] for n in nodes if "text" in n] == ["Before\tpage\nline", "After page"]
    for node in (nodes[0], nodes[2]):
        runs = node["attributes"]["richText"]
        assert all(r.get("marks") == [{"type": "bold"}] for r in runs if r["type"] == "text")
        assert "".join("\n" if r["type"] == "hardBreak" else r["text"] for r in runs) == node["text"]


def test_break_only_paragraphs_and_consecutive_manual_breaks_survive():
    doc = Document()
    doc.add_paragraph("First")
    doc.add_page_break()
    doc.add_page_break()
    doc.add_paragraph("Last")
    book, _ = parse_docx(archive(doc))
    assert [n["type"] for n in all_nodes(book)] == ["paragraph", "pageBreak", "pageBreak", "paragraph"]


def test_page_break_before_uses_style_inheritance_and_explicit_false():
    doc = Document()
    base = doc.styles.add_style("New sheet", WD_STYLE_TYPE.PARAGRAPH)
    base.paragraph_format.page_break_before = True
    child = doc.styles.add_style("Inherited sheet", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    doc.add_paragraph("First", style=child)  # Already at the chapter start.
    doc.add_paragraph("Second", style=child)
    doc.add_paragraph("Same page", style=child).paragraph_format.page_break_before = False
    doc.add_page_break()
    doc.add_paragraph("Already separated", style=child)
    book, _ = parse_docx(archive(doc))
    assert [n["type"] for n in all_nodes(book)] == [
        "paragraph", "pageBreak", "paragraph", "paragraph", "pageBreak", "paragraph"]


def test_custom_heading_styles_split_chapters_and_direct_outline_overrides():
    doc = Document()
    chapter = doc.styles.add_style("Novel chapter", WD_STYLE_TYPE.PARAGRAPH)
    chapter.base_style = doc.styles["Heading 1"]
    subheading = doc.styles.add_style("Scene label", WD_STYLE_TYPE.PARAGRAPH)
    subheading.base_style = doc.styles["Heading 2"]
    doc.add_paragraph("Opening", style=chapter)
    doc.add_paragraph("The harbor", style=subheading)
    body = doc.add_paragraph("A heading-styled body paragraph", style=chapter)
    outline(body._p.get_or_add_pPr(), 9)
    direct = doc.add_paragraph("Second chapter")
    outline(direct._p.get_or_add_pPr(), 0)
    book, _ = parse_docx(archive(doc))
    assert [c["title"] for c in book["chapters"]] == ["Opening", "Second chapter"]
    assert [(n["type"], n.get("level")) for n in book["chapters"][0]["nodes"]] == [
        ("heading", 1), ("heading", 2), ("paragraph", None)]
    assert book["chapters"][1]["nodes"][0]["level"] == 1


def test_heading_page_break_before_does_not_add_an_empty_chapter_page():
    doc = Document()
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    doc.add_heading("First chapter", 1)
    doc.add_paragraph("Body")
    doc.add_heading("Second chapter", 1)
    doc.add_paragraph("More")
    book, _ = parse_docx(archive(doc))
    assert len(book["chapters"]) == 2
    assert not any(n["type"] == "pageBreak" for n in all_nodes(book))


def test_outline_levels_beyond_the_editor_range_are_reported_and_bounded():
    doc = Document()
    outline(doc.add_paragraph("Deep heading")._p.get_or_add_pPr(), 8)
    book, report = parse_docx(archive(doc))
    assert all_nodes(book)[0]["level"] == 6
    assert any("heading" in warning.lower() and "six" in warning.lower() for warning in report["warnings"])


@pytest.mark.parametrize("value", ["-1", "10", "unknown", "0_1", "9" * 100])
def test_invalid_outline_levels_fail_with_an_actionable_import_error(value):
    doc = Document()
    outline(doc.add_paragraph("Body")._p.get_or_add_pPr(), value)
    with pytest.raises(ParseError, match="outline level"):
        parse_docx(archive(doc))


def test_cached_pagination_is_not_a_manual_page_break():
    doc = Document()
    run = doc.add_paragraph().add_run("Same")
    run._r.append(OxmlElement("w:lastRenderedPageBreak"))
    run.add_text(" page")
    book, _ = parse_docx(archive(doc))
    assert [(n["type"], n["text"]) for n in all_nodes(book)] == [("paragraph", "Same page")]


def test_column_break_is_reported_and_keeps_separated_text():
    doc = Document()
    run = doc.add_paragraph().add_run("Left")
    run.add_break(WD_BREAK.COLUMN)
    run.add_text("Right")
    book, report = parse_docx(archive(doc))
    assert all_nodes(book)[0]["text"] == "Left\nRight"
    assert any("column break" in warning.lower() for warning in report["warnings"])


def test_style_level_false_and_based_on_cycle_are_bounded():
    doc = Document()
    parent = doc.styles.add_style("Parent page", WD_STYLE_TYPE.PARAGRAPH)
    parent.paragraph_format.page_break_before = True
    child = doc.styles.add_style("Child no page", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = parent
    child.paragraph_format.page_break_before = False
    doc.add_paragraph("First")
    doc.add_paragraph("Still first page", style=child)
    a = doc.styles.add_style("Cycle A", WD_STYLE_TYPE.PARAGRAPH)
    b = doc.styles.add_style("Cycle B", WD_STYLE_TYPE.PARAGRAPH)
    a.base_style, b.base_style = b, a
    doc.add_paragraph("Cyclic style text", style=a)
    book, _ = parse_docx(archive(doc))
    assert [n["type"] for n in all_nodes(book)] == ["paragraph"] * 3
    assert all_nodes(book)[-1]["text"] == "Cyclic style text"


def test_hyperlink_page_break_keeps_labels_without_urls_or_field_instructions():
    doc = Document()
    paragraph = doc.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), "unresolved-external-link")
    run = OxmlElement("w:r")
    for tag, value in (("t", "First label"), ("br", None), ("instrText", "UNTRUSTED FIELD"), ("t", "Second label")):
        item = OxmlElement("w:" + tag)
        if tag == "br":
            item.set(qn("w:type"), "page")
        else:
            item.text = value
        run.append(item)
    link.append(run)
    paragraph._p.append(link)
    book, _ = parse_docx(archive(doc))
    nodes = all_nodes(book)
    assert [(n["type"], n.get("text")) for n in nodes] == [
        ("paragraph", "First label"), ("pageBreak", None), ("paragraph", "Second label")]
    assert "UNTRUSTED FIELD" not in str(book) and "unresolved-external-link" not in str(book)
