"""Real archive imports retain canonical text and supported formatting."""
import io
import json
import sys
from pathlib import Path

import pytest
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from parsers import ParseError
from parsers.docx_parser import parse_docx
from parsers.epub_parser import _xhtml_to_nodes, parse_epub


def flattened(n):
    return "".join("\n" if r["type"] == "hardBreak" else r["text"] for r in n["attributes"]["richText"])


def test_docx_real_archive_marks_hyperlink_labels_and_style_inheritance():
    doc = Document()
    doc.add_heading("Chapter One", 1)
    p = doc.add_paragraph()
    p.add_run("Hello ")
    p.add_run("bold").bold = True
    r = p.add_run("\nnext")
    r.italic = True
    r.underline = True
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), "never-resolved")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = " link label"
    run.append(text)
    link.append(run)
    p._p.append(link)
    doc.add_paragraph("First numbered", style="List Number")
    doc.add_paragraph("Nested bullet", style="List Bullet 2")
    p = doc.add_paragraph("Inherited italic", style="Quote")
    p.style.font.italic = True
    p.add_run(" plain").italic = False
    buf = io.BytesIO()
    doc.save(buf)
    book, report = parse_docx(buf.getvalue())
    nodes = book["chapters"][0]["nodes"]
    assert report["chapterCount"] == 1
    assert nodes[1]["text"] == "Hello bold\nnext link label"
    assert {"type": "bold"} in nodes[1]["attributes"]["richText"][1]["marks"]
    assert nodes[2]["attributes"]["listStyle"] == "ordered"
    assert nodes[3]["attributes"]["listStyle"] == "bullet"
    assert nodes[3]["attributes"]["listDepth"] == 1
    assert nodes[4]["attributes"]["richText"][0]["marks"] == [{"type": "italic"}]
    assert "marks" not in nodes[4]["attributes"]["richText"][1]
    assert all(flattened(n) == n["text"] for n in nodes)
    assert "never-resolved" not in json.dumps(book)


def test_epub_nested_blocks_once_no_punctuation_spaces_or_executable_markup():
    soup = BeautifulSoup('''<html><head><title>Not manuscript</title></head><body>
    <h1>Story</h1><p>Hello <strong>world</strong>, <em>reader</em>!<br/>Next.</p>
    <blockquote><p>One <b>quote</b>.</p><p>Second quote.</p></blockquote>
    <ol><li><p>Parent</p><p>continued</p><ul><li>Child <u>item</u></li></ul></li><li>Next item</li></ol>
    <p><a href="javascript:alert(1)">Safe label</a><script>BAD</script><span hidden>HIDDEN</span></p>
    <nav><p>Navigation</p></nav></body></html>''', "lxml")
    nodes = _xhtml_to_nodes(soup)
    assert [n["text"] for n in nodes] == ["Story", "Hello world, reader!\nNext.",
        "One quote.", "Second quote.", "Parent\ncontinued", "Child item", "Next item", "Safe label"]
    assert [n["type"] for n in nodes[2:4]] == ["quote", "quote"]
    assert [(n["attributes"]["listStyle"], n["attributes"]["listDepth"]) for n in nodes[4:7]] == [
        ("ordered", 0), ("bullet", 1), ("ordered", 0)]
    assert all(flattened(n) == n["text"] for n in nodes)
    saved = json.dumps(nodes)
    assert all(s not in saved for s in ("javascript:", "BAD", "HIDDEN", "Navigation"))


def test_epub_nesting_limit_rejects_instead_of_recursing_unbounded():
    with pytest.raises(ParseError, match="nesting"):
        _xhtml_to_nodes(BeautifulSoup("<div>" * 140 + "Text" + "</div>" * 140, "lxml"))


def test_import_export_import_keeps_quotes_lists_marks_and_text_once():
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "rendering"))
    from editions import EbookEdition
    from epub_renderer import render_epub
    from parsers import make_book, new_chapter
    chapter = new_chapter("Sample", 0)
    chapter["nodes"] = _xhtml_to_nodes(BeautifulSoup(
        '<h1>Sample</h1><blockquote><p>A <b>quote</b>.</p></blockquote>'
        '<ol><li>Parent<ul><li><em>Child</em></li></ul></li></ol>', "lxml"))
    book = make_book([chapter])
    rendered = render_epub(book, EbookEdition())
    data = rendered[0] if isinstance(rendered, tuple) else rendered
    imported, report = parse_epub(data)
    def semantic(nodes):
        return [{k: v for k, v in n.items() if k != "id"} for n in nodes]
    assert report["chapterCount"] == 1
    assert semantic(imported["chapters"][0]["nodes"]) == semantic(chapter["nodes"])


def test_docx_tables_stay_between_paragraphs_and_reach_epub_and_pdf():
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "rendering"))
    from editions import EbookEdition, PrintEdition
    from epub_renderer import render_epub
    from pdf_renderer import render_pdf
    from pypdf import PdfReader
    doc = Document()
    doc.add_paragraph("Before table.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Description"
    table.cell(1, 0).text = "Fox & Finch"
    table.cell(1, 1).text = "First line\nSecond line <literal>"
    doc.add_paragraph("After table.")
    buf = io.BytesIO()
    doc.save(buf)
    book, report = parse_docx(buf.getvalue())
    nodes = book["chapters"][0]["nodes"]
    assert [n["type"] for n in nodes] == ["paragraph", "table", "paragraph"]
    assert nodes[1]["rows"] == [["Name", "Description"], ["Fox & Finch", "First line\nSecond line <literal>"]]
    assert any("merged-cell" in warning for warning in report["warnings"])
    epub, _ = render_epub(book, EbookEdition())
    again, _ = parse_epub(epub)
    assert again["chapters"][0]["nodes"][1]["rows"] == nodes[1]["rows"]
    pdf, checksum = render_pdf(book, PrintEdition())
    assert render_pdf(book, PrintEdition())[1] == checksum
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert text.index("Before table.") < text.index("Fox & Finch") < text.index("After table.")
    assert "Second line <literal>" in text


def test_docx_merged_cells_do_not_duplicate_source_text():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 1)).text = "Only once"
    buf = io.BytesIO()
    doc.save(buf)
    book, _ = parse_docx(buf.getvalue())
    assert book["chapters"][0]["nodes"][0]["text"].count("Only once") == 1
