import base64
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import sys
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient

SERVICE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SERVICE))
from parsers import ParseError
from parsers.docx_parser import parse_docx
from parsers.epub_parser import parse_epub
from parsers.embedded_images import EmbeddedImages

PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=")


def epub_with_images():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("META-INF/container.xml", '<container><rootfiles><rootfile full-path="OPS/content.opf"/></rootfiles></container>')
        z.writestr("OPS/content.opf", '<package xmlns="http://www.idpf.org/2007/opf"><manifest><item id="ch" href="chapter.xhtml"/></manifest><spine><itemref idref="ch"/></spine></package>')
        z.writestr("OPS/chapter.xhtml", '<html><body><p>Before</p><img src="images/a%20b.png" alt="Harbor"/>'
            '<p>Between</p><img src="images/a%20b.png" alt=""/><p>After</p>'
            '<img src="https://external.invalid/image.png" alt="External reference"/>'
            '<img src="../../missing.png" alt="Missing reference"/></body></html>')
        z.writestr("OPS/images/a b.png", PNG)
    return buf.getvalue()


def test_docx_inline_images_preserve_text_order_marks_alt_and_deduplicated_bytes():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Before").bold = True
    picture = p.add_run().add_picture(io.BytesIO(PNG))
    picture._inline.docPr.set("descr", "Harbor at dawn")
    p.add_run("Between").italic = True
    p.add_run().add_picture(io.BytesIO(PNG))
    p.add_run("After")
    buf = io.BytesIO(); doc.save(buf)
    images = []
    book, report = parse_docx(buf.getvalue(), embedded_assets=images)
    nodes = book["chapters"][0]["nodes"]
    assert [n["type"] for n in nodes] == ["paragraph", "image", "paragraph", "image", "paragraph"]
    assert [n["text"] for n in nodes if "text" in n] == ["Before", "Between", "After"]
    assert nodes[0]["attributes"]["richText"][0]["marks"] == [{"type": "bold"}]
    assert nodes[1]["altText"] == "Harbor at dawn"
    assert nodes[1]["assetId"] == nodes[3]["assetId"] == images[0]["id"]
    assert len(images) == report["imageCount"] == 1
    assert base64.b64decode(images[0]["contentBase64"]) == PNG
    assert images[0]["checksumSha256"] == hashlib.sha256(PNG).hexdigest()
    assert "contentBase64" not in json.dumps(book)
    assert "contentBase64" not in json.dumps(report)


def test_epub_local_images_keep_placement_and_external_paths_are_never_transferred():
    images = []
    book, report = parse_epub(epub_with_images(), embedded_assets=images)
    nodes = book["chapters"][0]["nodes"]
    assert [n.get("text") for n in nodes[:5]] == ["Before", None, "Between", None, "After"]
    assert nodes[1]["assetId"] == nodes[3]["assetId"] == images[0]["id"]
    assert nodes[1]["altText"] == "Harbor"
    assert nodes[3]["attributes"]["decorative"] is True
    assert len(images) == 1
    assert report["warnings"] and report["imageCount"] == 1
    assert "https://" not in json.dumps(book) and "../" not in json.dumps(book)


def test_authenticated_epub_endpoint_accepts_collector_and_returns_private_payload_separately(monkeypatch):
    spec = importlib.util.spec_from_file_location("document_embedded_http", SERVICE / "main.py")
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
    monkeypatch.setenv("DOCUMENT_SERVICE_TOKEN", "fixture-embedded-token")
    with TestClient(module.app) as client:
        response = client.post("/parse", headers={"x-service-token": "fixture-embedded-token"}, json={
            "assetId": "source", "format": "epub", "contentBase64": base64.b64encode(epub_with_images()).decode()})
    assert response.status_code == 200, response.text
    payload = response.json()
    assert len(payload["embeddedAssets"]) == 1
    assert "contentBase64" not in json.dumps(payload["report"])
    assert "contentBase64" not in json.dumps(payload["bookModel"])


def test_image_collector_limits_and_unsupported_types(monkeypatch):
    import parsers.embedded_images as images
    collector = EmbeddedImages()
    assert collector.add(b"<svg/>", "image/svg+xml") is None
    monkeypatch.setattr(images, "MAX_IMAGE_BYTES", 2)
    with pytest.raises(ParseError, match="10 MiB"):
        collector.add(PNG, "image/png")
    monkeypatch.setattr(images, "MAX_IMAGE_BYTES", 1024)
    monkeypatch.setattr(images, "MAX_IMAGES", 1)
    first = collector.add(PNG, "image/png")
    assert collector.add(PNG, "image/png") == first
    with pytest.raises(ParseError, match="budget"):
        collector.add(PNG + b"different", "image/png")
