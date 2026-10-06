"""Real Word numbering survives the canonical model and deterministic exports."""
import io
import sys
from pathlib import Path
from zipfile import ZipFile

import pytest
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "rendering"))
from parsers import ParseError
from parsers.docx_parser import parse_docx
from parsers.epub_parser import parse_epub, _xhtml_to_nodes
from editions import EbookEdition, PrintEdition
from epub_renderer import render_epub
from pdf_renderer import render_pdf
from pypdf import PdfReader


def child(parent, name, value=None, **attrs):
    element = OxmlElement("w:" + name)
    if value is not None:
        element.set(qn("w:val"), str(value))
    for key, val in attrs.items():
        element.set(qn("w:" + key), str(val))
    parent.append(element)
    return element


def numbered_document():
    doc = Document()
    abstract = child(doc.part.numbering_part.element, "abstractNum", abstractNumId=700)
    for level, fmt in enumerate(("lowerRoman", "upperLetter")):
        definition = child(abstract, "lvl", ilvl=level)
        child(definition, "start", 1)
        child(definition, "numFmt", fmt)
        child(definition, "lvlText", "%" + str(level + 1) + ".")
    for num_id, start in ((700, 7), (701, 2)):
        num = child(doc.part.numbering_part.element, "num", numId=num_id)
        child(num, "abstractNumId", 700)
        override = child(num, "lvlOverride", ilvl=0)
        child(override, "startOverride", start)
    return doc


def number(properties, num_id, level=0):
    num = child(properties, "numPr")
    child(num, "numId", num_id)
    child(num, "ilvl", level)


def source_bytes(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_word_direct_and_inherited_numbering_restart_continuation_and_nesting():
    doc = numbered_document()
    style = doc.styles.add_style("Author Steps", WD_STYLE_TYPE.PARAGRAPH)
    number(style.element.get_or_add_pPr(), 700)
    derived = doc.styles.add_style("Author Steps Child", WD_STYLE_TYPE.PARAGRAPH)
    derived.base_style = style
    doc.add_paragraph("First step", style=derived)
    number(doc.add_paragraph("Nested step")._p.get_or_add_pPr(), 700, 1)
    doc.add_paragraph("Second step", style=derived)
    doc.add_paragraph("Intervening text")
    doc.add_paragraph("Continued step", style=derived)
    number(doc.add_paragraph("Restarted step")._p.get_or_add_pPr(), 701)
    suppressed = doc.add_paragraph("Not a numbered step", style=derived)
    number(suppressed._p.get_or_add_pPr(), 0)
    book, _ = parse_docx(source_bytes(doc))
    nodes = book["chapters"][0]["nodes"]
    assert [n["type"] for n in nodes] == ["listItem", "listItem", "listItem", "paragraph", "listItem", "listItem", "paragraph"]
    assert [(n["attributes"].get("listStart"), n["attributes"].get("listNumberStyle"), n["attributes"].get("listDepth"))
            for n in nodes if n["type"] == "listItem"] == [(7, "lower-roman", 0), (1, "upper-alpha", 1), (8, "lower-roman", 0), (9, "lower-roman", 0), (2, "lower-roman", 0)]


def test_epub_starts_item_values_and_styles_are_preserved_without_executable_attributes():
    nodes = _xhtml_to_nodes(BeautifulSoup('<ol start="7" type="i"><li>First</li><li value="12">Jump</li>'
        '<li>Next<ol type="A" start="3"><li>Nested</li></ol></li></ol>', "lxml"))
    assert [n["attributes"].get("listStart") for n in nodes] == [7, 12, 13, 3]
    assert [n["attributes"].get("listNumberStyle") for n in nodes] == ["lower-roman"] * 3 + ["upper-alpha"]
    malicious = _xhtml_to_nodes(BeautifulSoup('<ol start="999999999999999999999999" type="script"><li value="javascript:alert(1)">Safe</li></ol>', "lxml"))
    assert malicious[0]["attributes"]["listStart"] == 1
    assert malicious[0]["attributes"]["listNumberStyle"] == "decimal"


def test_word_numbering_survives_real_epub_reimport_and_pdf_markers():
    doc = numbered_document()
    number(doc.add_paragraph("First step")._p.get_or_add_pPr(), 700)
    number(doc.add_paragraph("Second step")._p.get_or_add_pPr(), 700)
    doc.add_paragraph("Interlude")
    number(doc.add_paragraph("Continued step")._p.get_or_add_pPr(), 700)
    number(doc.add_paragraph("Restarted step")._p.get_or_add_pPr(), 701)
    book, _ = parse_docx(source_bytes(doc))
    epub, digest = render_epub(book, EbookEdition())
    assert render_epub(book, EbookEdition())[1] == digest
    with ZipFile(io.BytesIO(epub)) as archive:
        markup = archive.read("OEBPS/ch0000.xhtml")
    assert b'<ol start="7" type="i">' in markup
    assert b'<ol start="9" type="i">' in markup
    assert b'<ol start="2" type="i">' in markup
    imported, _ = parse_epub(epub)
    assert [(n["text"], n["attributes"].get("listStart"), n["attributes"].get("listNumberStyle")) for n in imported["chapters"][0]["nodes"]] == [
        (n["text"], n["attributes"].get("listStart"), n["attributes"].get("listNumberStyle")) for n in book["chapters"][0]["nodes"]]
    pdf, checksum = render_pdf(book, PrintEdition())
    assert render_pdf(book, PrintEdition())[1] == checksum
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert all(marker in text for marker in ("vii", "viii", "ix", "ii"))


def test_word_unsupported_compound_numbering_is_reported_not_claimed_exact():
    doc = numbered_document()
    abstract = doc.part.numbering_part.element.find('w:abstractNum[@w:abstractNumId="700"]', doc.part.numbering_part.element.nsmap)
    abstract.find("w:lvl/w:lvlText", abstract.nsmap).set(qn("w:val"), "%1.%2.")
    number(doc.add_paragraph("Outline")._p.get_or_add_pPr(), 700)
    _, report = parse_docx(source_bytes(doc))
    assert any("compound" in warning for warning in report["warnings"])


def test_alphabetic_print_does_not_wrap_z_to_a_and_long_markers_leave_space():
    doc = numbered_document()
    for text in ("Before rollover", "After rollover"):
        number(doc.add_paragraph(text)._p.get_or_add_pPr(), 700, 1)
    book, _ = parse_docx(source_bytes(doc))
    for index, node in enumerate(book["chapters"][0]["nodes"]):
        node["attributes"]["listStart"] = 26 + index
    pdf, _ = render_pdf(book, PrintEdition())
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "Z." in text and "AA." in text


@pytest.mark.parametrize("start", ["0", "-1", "1000001", "999999999999999999999999", "text"])
def test_invalid_word_start_fails_safely_without_rewriting_the_number(start):
    doc = numbered_document()
    num = doc.part.numbering_part.element.find('w:num[@w:numId="700"]', doc.part.numbering_part.element.nsmap)
    num.find("w:lvlOverride/w:startOverride", num.nsmap).set(qn("w:val"), start)
    number(doc.add_paragraph("Unsafe numbering")._p.get_or_add_pPr(), 700)
    with pytest.raises(ParseError, match="list start"):
        parse_docx(source_bytes(doc))


def test_epub_counter_overflow_fails_instead_of_repeating_the_last_number():
    with pytest.raises(ParseError, match="numbering exceeds"):
        _xhtml_to_nodes(BeautifulSoup('<ol start="1000000"><li>Last</li><li>Overflow</li></ol>', "lxml"))


def test_reversed_epub_lists_retain_direction_starts_and_custom_item_values_in_print():
    nodes = _xhtml_to_nodes(BeautifulSoup('<ol reversed="reversed" type="A"><li>First</li><li value="8">Second</li><li>Third</li></ol>', "lxml"))
    assert [n["attributes"]["listStart"] for n in nodes] == [3, 8, 7]
    assert all(n["attributes"]["listReversed"] is True for n in nodes)
    doc = Document(); doc.add_paragraph("Placeholder")
    book, _ = parse_docx(source_bytes(doc)); book["chapters"][0]["nodes"] = nodes
    epub, _ = render_epub(book, EbookEdition())
    again, _ = parse_epub(epub)
    assert [(n["attributes"]["listStart"], n["attributes"]["listReversed"]) for n in again["chapters"][0]["nodes"]] == [(3, True), (8, True), (7, True)]
    pdf, _ = render_pdf(book, PrintEdition())
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
    assert all(value in text for value in ("C.", "H.", "G."))
