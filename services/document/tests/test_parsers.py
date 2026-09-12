"""Parser tests: deterministic Book Model shape, corrupt input -> ParseError, zip-slip rejected."""
import io
import zipfile

import pytest

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from parsers import ParseError, safe_zip_members
from parsers.docx_parser import parse_docx
from parsers.epub_parser import parse_epub
from parsers.pdf_parser import parse_pdf
from parsers.txt_parser import parse_txt

NODE_TYPES = {"paragraph", "heading", "quote", "list", "listItem", "image",
              "caption", "pageBreak", "table", "footnote", "separator"}


def make_docx() -> bytes:
    from docx import Document
    doc = Document()
    doc.add_heading("Chapter One", level=1)
    doc.add_paragraph("Hello world.")
    doc.add_paragraph("Second para.")
    doc.add_heading("Chapter Two", level=1)
    doc.add_paragraph("More text.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_epub(slip: bool = False) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        if slip:
            z.writestr("../../evil.txt", "pwned")
        z.writestr("META-INF/container.xml", """<?xml version="1.0"?>
<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">
  <rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>""")
        z.writestr("OEBPS/content.opf", """<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:title id="id">Fixture Book</dc:title>
  </metadata>
  <manifest>
    <item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/>
    <item id="c2" href="ch2.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine><itemref idref="c1"/><itemref idref="c2"/></spine>
</package>""")
        for i in (1, 2):
            z.writestr(f"OEBPS/ch{i}.xhtml", f"""<?xml version="1.0"?>
<html xmlns="http://www.w3.org/1999/xhtml"><body>
<h1>Chapter {i}</h1><p>Paragraph one of chapter {i}.</p><p>Second paragraph.</p>
</body></html>""")
    return buf.getvalue()


def assert_book_shape(book: dict, report: dict, min_chapters: int = 1) -> None:
    assert book["schemaVersion"] == "1.0"
    assert book["metadata"]["title"]
    assert book["metadata"]["author"]
    assert len(book["chapters"]) >= min_chapters
    for ch in book["chapters"]:
        assert {"id", "order", "title", "nodes"} <= set(ch)
        for n in ch["nodes"]:
            assert n["type"] in NODE_TYPES
            assert n["id"]
    assert report["chapterCount"] == len(book["chapters"])
    assert report["nodeCount"] == sum(len(c["nodes"]) for c in book["chapters"])
    assert isinstance(report["warnings"], list)


def test_docx_two_chapters():
    book, report = parse_docx(make_docx(), title="D")
    assert_book_shape(book, report, 2)
    assert [c["title"] for c in book["chapters"]] == ["Chapter One", "Chapter Two"]
    assert book["chapters"][0]["nodes"][0] == {
        **book["chapters"][0]["nodes"][0], "type": "heading", "level": 1}


def test_docx_corrupt():
    with pytest.raises(ParseError):
        parse_docx(b"not a docx at all")


def test_epub_spine_order():
    book, report = parse_epub(make_epub())
    assert_book_shape(book, report, 2)
    assert book["metadata"]["title"] == "Fixture Book"
    assert [c["title"] for c in book["chapters"]] == ["Chapter 1", "Chapter 2"]
    assert book["chapters"][0]["nodes"][0]["type"] == "heading"


def test_epub_zip_slip_rejected():
    with pytest.raises(ParseError, match="zip-slip|unsafe"):
        parse_epub(make_epub(slip=True))


def test_epub_corrupt():
    with pytest.raises(ParseError):
        parse_epub(b"\x00\x01\x02")


def test_safe_zip_members_direct():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("a/../b.txt", "x")
    buf.seek(0)
    with zipfile.ZipFile(buf) as zf:
        with pytest.raises(ParseError):
            safe_zip_members(zf)


def test_txt_chapter_heuristics():
    text = ("Chapter 1\n\nFirst para here.\n\n"
            "Chapter 2\n\nAnother para.\n\nStill chapter two.\n")
    book, report = parse_txt(text.encode("utf-8"), title="T")
    assert_book_shape(book, report, 2)
    assert [c["title"] for c in book["chapters"]] == ["Chapter 1", "Chapter 2"]
    assert len(book["chapters"][1]["nodes"]) == 3  # heading + 2 paras


def test_txt_allcaps_and_fallback():
    book, _ = parse_txt(b"PROLOGUE\n\nSome words.\n", title="T")
    assert book["chapters"][0]["title"] == "PROLOGUE"
    book2, _ = parse_txt("Just one long block\nof plain text.\n".encode(), title="T")
    assert_book_shape(book2, _, 1)


def test_txt_encoding_detection():
    book, report = parse_txt("Café déjà vu.".encode("cp1252"), title="T")
    assert_book_shape(book, report, 1)
    assert "Caf" in book["chapters"][0]["nodes"][0]["text"]


def test_txt_binary_rejected():
    with pytest.raises(ParseError):
        parse_txt(b"\x00\x00\x00\x00binary")


def test_pdf_garbage_rejected():
    with pytest.raises(ParseError):
        parse_pdf(b"definitely not a pdf")


def test_pdf_low_confidence_path():
    # pypdf writer can't embed extractable text reliably; build a minimal PDF with
    # no text objects so extraction is empty -> sparse -> low confidence + QA warning.
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    w.write(buf)
    book, report = parse_pdf(buf.getvalue(), title="P")
    assert_book_shape(book, report, 1)
    assert report["confidence"] == "low"
    assert any("manual QA" in w_ or "OCR" in w_ for w_ in report["warnings"])


def test_http_422_on_corrupt(monkeypatch):
    monkeypatch.setenv("DOCUMENT_SERVICE_TOKEN", "fixture-parser-token")
    from fastapi.testclient import TestClient
    import base64
    # Drop any cached `main` (collides across services in combined pytest runs)
    # so this service's own main.py is imported fresh.
    sys.modules.pop("main", None)
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    from main import app
    client = TestClient(app)
    r = client.post("/parse", headers={"x-service-token": "fixture-parser-token"}, json={
        "assetId": "a1", "format": "epub",
        "contentBase64": base64.b64encode(b"junk").decode()})
    assert r.status_code == 422
    ok = client.post("/parse", headers={"x-service-token": "fixture-parser-token"}, json={
        "assetId": "a2", "format": "txt",
        "contentBase64": base64.b64encode("Chapter 1\n\nHi.\n".encode()).decode()})
    assert ok.status_code == 200
    assert_book_shape(ok.json()["bookModel"], ok.json()["report"], 1)
